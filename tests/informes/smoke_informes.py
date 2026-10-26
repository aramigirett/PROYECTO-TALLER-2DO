"""Módulo Informes (pantalla única /informes, solo Administrador, solo consulta).
Etapa 1 y 2: migración 03 (idempotente, vistas, permisos, CHECK 'CONSULTA', menú), acceso (solo Administrador), filtros, auditoría de visitas y
consultas, y el contenido de las cinco pestañas (Resumen, Agendamiento, Consultorio, Facturación, Auditoría) contrastado con SQL independiente.
Etapa 3: exportación a PDF y Excel por pestaña con su registro en auditoría.
Datos de prueba: citas / consulta del año 2001 (ventana que no usa ningún dato real), creadas y borradas POR ID. La auditoría que genera el test
se borra por ID (saltando el trigger solo en su conexión)."""
import io, pathlib, sys
RAIZ = pathlib.Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RAIZ))
from app import app
from tests._conexion_pruebas import Conexion
from app.rutas.modulo_seguridad.menu_catalog import CATALOGO_MENU, VENTANAS_SOLO_ADMIN

fallos = []
c = app.test_client()
V03 = RAIZ / 'migraciones' / '03_migracion_informes.sql'
VISTAS = ['v_inf_citas', 'v_inf_consultas', 'v_inf_diagnosticos', 'v_inf_tratamientos', 'v_inf_procedimientos', 'v_inf_insumos',
          'v_inf_comprobantes', 'v_inf_cobros', 'v_inf_cuentas_cobrar', 'v_inf_accesos']
W = 'desde=2001-03-01&hasta=2001-03-31'          # ventana de los datos de prueba
AMPLIA = 'desde=2020-01-01&hasta=2026-12-31'


def chk(nombre, cond):
    print(('OK  ' if cond else 'FAIL'), nombre)
    if not cond:
        fallos.append(nombre)


cx = Conexion().getConexion(); cx.autocommit = True; cu = cx.cursor()
cu.execute("SELECT COALESCE(MAX(id_auditoria_cambio), 0) FROM auditoria_cambios"); BASE = cu.fetchone()[0]
cu.execute("SELECT u.id_usuario, r.id_rol, r.nombre_rol FROM usuarios u JOIN roles r USING(id_rol) WHERE u.activo ORDER BY 1")
usuarios = cu.fetchall()
ADMIN = next((u for u in usuarios if u[2] == 'Administrador'), None)
CAJERO = next((u for u in usuarios if u[2] == 'Cajero'), None)
if not (ADMIN and CAJERO):
    print('ABORTADO: se necesitan un Administrador y un Cajero.'); sys.exit(1)


def como(u):
    with c.session_transaction() as s:
        s.clear()
        if u:
            s['id_usuario'] = u[0]; s['id_rol'] = u[1]; s['nombre_rol'] = u[2]; s['nombre_completo'] = 'USUARIO PRUEBA'


def api(pestana, filtros=W):
    r = c.get(f'/api/informes/{pestana}?{filtros}')
    return r, (r.get_json(silent=True) or {})


def kpi(d, clave):
    return next(k['valor'] for k in d['data']['kpis'] if k['clave'] == clave)


def tabla(d, id_):
    return next(t for t in d['data']['tablas'] if t['id'] == id_)


def uno(sql, params=()):
    cu.execute(sql, params)
    return cu.fetchone()[0]


citas_zz, consultas_zz = [], []
try:
    # =========================================================================================================================
    # 1. MIGRACIÓN 03: idempotente, vistas, permisos, CHECK, menú
    # =========================================================================================================================
    cx.set_client_encoding('UTF8')
    cu.execute(V03.read_text(encoding='utf-8'))      # segunda aplicación: no debe fallar ni duplicar
    cu.execute(V03.read_text(encoding='utf-8'))
    chk('migración 03 idempotente (se aplicó dos veces sin error)', True)
    cu.execute("SELECT table_name FROM information_schema.views WHERE table_schema = 'public' AND table_name LIKE 'v\\_inf\\_%'")
    chk('las 10 vistas v_inf_* existen', {r[0] for r in cu.fetchall()} == set(VISTAS))
    chk('odontoclinic_app solo puede SELECT en las vistas (sin INSERT / UPDATE / DELETE)', all(
        uno("SELECT has_table_privilege('odontoclinic_app', %s, 'INSERT') OR has_table_privilege('odontoclinic_app', %s, 'UPDATE') OR has_table_privilege('odontoclinic_app', %s, 'DELETE')", (v, v, v)) is False for v in VISTAS))
    chk('odontoclinic_app puede SELECT en todas las vistas',
        all(uno("SELECT has_table_privilege('odontoclinic_app', %s, 'SELECT')", (v,)) for v in VISTAS))
    cu.execute("SELECT pg_get_constraintdef(oid) FROM pg_constraint WHERE conrelid = 'auditoria_cambios'::regclass AND conname = 'auditoria_cambios_accion_check'")
    definicion = cu.fetchone()[0]
    chk("el CHECK de acción admite CONSULTA y conserva las anteriores", all(a in definicion for a in ('CONSULTA', 'EXPORTACION', 'CIERRE', 'ALTA')))
    chk("una sola fila de menú 'informes' (activa, solo_admin, bajo el módulo Informes)",
        uno("""SELECT count(*) FROM menu v JOIN menu p ON p.id_menu = v.id_menu_padre WHERE v.clave = 'informes' AND v.activo AND v.solo_admin
               AND v.ruta = 'informes.informesIndex' AND p.descripcion = 'Informes' AND p.clave IS NULL""") == 1
        and uno("SELECT count(*) FROM menu WHERE clave IS NULL AND id_menu_padre IS NULL AND UPPER(descripcion) = 'INFORMES'") == 1)
    chk("el catálogo del código declara 'informes' como solo_admin", 'informes' in VENTANAS_SOLO_ADMIN
        and any(i['ventana'] == 'informes' and i['modulo'] == 'Informes' for i in CATALOGO_MENU))

    # =========================================================================================================================
    # 2. ACCESO: solo Administrador
    # =========================================================================================================================
    como(None)
    r = c.get('/informes')
    chk('sin sesión: la vista redirige al login y el API responde 401', r.status_code in (301, 302) and c.get(f'/api/informes/resumen?{W}').status_code == 401)
    como(CAJERO)
    chk('Cajero: vista 403 y API 403 (pestañas, médicos y exportación)', c.get('/informes').status_code == 403 and c.get(f'/api/informes/resumen?{W}').status_code == 403
        and c.get('/api/informes/medicos').status_code == 403 and c.get(f'/api/informes/facturacion/exportar?formato=pdf&{W}').status_code in (403, 404))
    como(ADMIN)
    r = c.get('/informes')
    html = r.data.decode('utf-8')
    chk('Administrador: la vista carga con pestañas, filtros y botones', r.status_code == 200 and all(t in html for t in
        ('data-pestana="resumen"', 'data-pestana="agendamiento"', 'data-pestana="consultorio"', 'data-pestana="facturacion"', 'data-pestana="auditoria"',
         'id="fDesde"', 'id="fHasta"', 'id="fMedico"', 'id="fCondicion"', 'id="btnAplicar"', 'id="btnPdf"', 'id="btnExcel"')))
    chk('el menú lateral del Administrador ofrece Informes', 'href="/informes"' in html)
    chk('ApexCharts y informes.js se sirven', c.get('/static/vendor/apexcharts/apexcharts.min.js').status_code == 200 and c.get('/static/js/informes.js').status_code == 200)
    cu.execute("SELECT detalle, modulo, id_usuario, ip FROM auditoria_cambios WHERE id_auditoria_cambio > %s AND entidad = 'informes' AND accion = 'CONSULTA' ORDER BY 1", (BASE,))
    visitas = cu.fetchall()
    chk('la visita a /informes queda auditada (CONSULTA, módulo Informes, usuario e IP del servidor)',
        len(visitas) >= 1 and visitas[0][0] == 'Visita a la pantalla de Informes' and visitas[0][1] == 'Módulo Informes' and visitas[0][2] == ADMIN[0] and visitas[0][3] == '127.0.0.1')

    # =========================================================================================================================
    # 3. FILTROS
    # =========================================================================================================================
    r, d = api('resumen', 'desde=2026-10-10&hasta=2026-10-01')
    chk('Desde > Hasta -> 400', r.status_code == 400 and not d['success'])
    chk('fecha inválida -> 400', api('resumen', 'desde=ayer')[0].status_code == 400)
    chk('id_medico inválido -> 400', api('agendamiento', f'{W}&id_medico=abc')[0].status_code == 400)
    chk('condición inválida -> 400', api('facturacion', f'{W}&condicion=otra')[0].status_code == 400)
    chk('pestaña inexistente -> 404', api('inexistente')[0].status_code == 404)
    r, d = api('resumen', '')
    chk('sin fechas: por defecto del día 1 del mes actual a hoy', r.status_code == 200 and d['success'])
    cu.execute("SELECT detalle FROM auditoria_cambios WHERE id_auditoria_cambio > %s AND entidad = 'informes' AND detalle LIKE 'Informe: Resumen%%' ORDER BY id_auditoria_cambio DESC LIMIT 1", (BASE,))
    ult = cu.fetchone()[0]
    chk('la auditoría de la consulta lleva el informe y el período por defecto', ult.startswith('Informe: Resumen; Período: 01/') and ' al ' in ult)
    como(ADMIN)
    r, d = api('resumen', 'desde=2020-01-01&hasta=2020-01-31&id_medico=1&condicion=CONTADO')
    chk('un filtro que no aplica a la pestaña se ignora (Resumen con condición)', r.status_code == 200)

    # =========================================================================================================================
    # 4. DATOS DE PRUEBA 2001: citas en cada categoría + una consulta
    # =========================================================================================================================
    cu.execute("SELECT id_cita_cabecera, id_agenda_detalle, hora_cita FROM cita_detalle WHERE id_cita_cabecera IN (SELECT id_cita_cabecera FROM cita_cabecera WHERE estado = 'Activo') ORDER BY 1 LIMIT 1")
    cab, agd, hora = cu.fetchone()
    cu.execute("""SELECT ac.id_medico, cc.id_paciente FROM cita_cabecera cc JOIN agenda_cabecera ac USING (id_agenda_cabecera) WHERE cc.id_cita_cabecera = %s""", (cab,))
    id_medico, id_paciente = cu.fetchone()

    def cita(estado, dia):
        cu.execute("""INSERT INTO cita_detalle (id_cita_cabecera, id_agenda_detalle, fecha_cita, hora_cita, motivo_consulta, id_estado_cita)
                      VALUES (%s, %s, %s, %s, 'ZZ INFORMES', %s) RETURNING id_cita_detalle""", (cab, agd, f'2001-03-{dia:02d}', hora, estado))
        i = cu.fetchone()[0]; citas_zz.append(i)
        return i

    c_aus1 = cita(2, 5)        # Confirmado, fecha pasada, sin consulta  -> AUSENTE
    c_conf = cita(2, 6)        # Confirmado con consulta                  -> no ausente
    cita(5, 7)                 # No Asistió                               -> AUSENTE
    cita(6, 8)                 # Reprogramado                             -> solo agendada
    cita(4, 9)                 # Cancelado
    cita(3, 10)                # Realizado
    cu.execute("""INSERT INTO consultas_cab (id_cita, id_paciente, id_medico, id_consultorio, id_funcionario, fecha_cita, hora_cita, duracion_minutos, estado, activo)
                  SELECT %s, id_paciente, id_medico, id_consultorio, id_funcionario, '2001-03-06', hora_cita, duracion_minutos, 'programada', true
                    FROM consultas_cab ORDER BY id_consulta_cab LIMIT 1 RETURNING id_consulta_cab, id_paciente, id_medico""", (c_conf,))
    cons, pac_cons, med_cons = cu.fetchone(); consultas_zz.append(cons)

    # =========================================================================================================================
    # 5. AGENDAMIENTO
    # =========================================================================================================================
    r, d = api('agendamiento')
    chk('Agendamiento: 200 con kpis, gráfico y 4 tablas', r.status_code == 200 and d['success'] and len(d['data']['graficos']) == 1
        and [t['id'] for t in d['data']['tablas']] == ['por_medico', 'ausentes', 'reiteradas', 'canceladas'])
    print('   KPIs:', {k['clave']: k['valor'] for k in d['data']['kpis']})
    chk('agendadas 6 (todas, Reprogramado incluido)', kpi(d, 'agendadas') == 6)
    chk('atendidas 1, canceladas 1', kpi(d, 'atendidas') == 1 and kpi(d, 'canceladas') == 1)
    chk('ausentes 2 (No Asistió + Confirmado pasada sin consulta, sin duplicar; el Confirmado con consulta no cuenta)', kpi(d, 'ausentes') == 2)
    chk('Reprogramado no es ausente ni cancelada: solo suma en agendadas', kpi(d, 'reprogramadas') == 1)
    chk('tasas: asistencia 33,3 % (1 / (1 + 2)), ausentismo 33,3 %, cancelación 16,7 %',
        kpi(d, 'tasa_asistencia') == 33.3 and kpi(d, 'tasa_ausentismo') == 33.3 and kpi(d, 'tasa_cancelacion') == 16.7)
    ev = d['data']['graficos'][0]
    chk('evolución diaria: 6 días, serie agendadas suma 6 y ausentes 2',
        len(ev['categorias']) == 6 and sum(next(s for s in ev['series'] if s['nombre'] == 'Agendadas')['datos']) == 6
        and sum(next(s for s in ev['series'] if s['nombre'] == 'Ausentes')['datos']) == 2)
    pm = tabla(d, 'por_medico')['filas']
    chk('por médico: una fila con 6 / 1 / 1 / 2 y asistencia 33,3', len(pm) == 1 and (pm[0]['agendadas'], pm[0]['atendidas'], pm[0]['canceladas'], pm[0]['ausentes'], pm[0]['tasa']) == (6, 1, 1, 2, 33.3))
    aus = tabla(d, 'ausentes')['filas']
    chk('pacientes ausentes: 2 filas con ausencias previas 1 (la más reciente) y 0 (la primera)', len(aus) == 2 and aus[0]['previas'] == 1 and aus[1]['previas'] == 0
        and aus[0]['fecha'] == '07/03/2001' and aus[1]['fecha'] == '05/03/2001')
    reit = tabla(d, 'reiteradas')['filas']
    chk('ausencias reiteradas (2+): el paciente con 2 ausencias', len(reit) == 1 and reit[0]['ausencias'] == 2 and reit[0]['ultima'] == '07/03/2001')
    can = tabla(d, 'canceladas')['filas']
    chk('citas canceladas: 1 fila (09/03/2001)', len(can) == 1 and can[0]['fecha'] == '09/03/2001')
    r, d2 = api('agendamiento', f'{W}&id_medico={id_medico}')
    chk('filtro por el médico de las citas: mismas 6', kpi(d2, 'agendadas') == 6)
    otro = uno("SELECT COALESCE(MIN(id_medico), 0) FROM medico WHERE id_medico <> %s AND activo", (id_medico,))
    if otro:
        r, d3 = api('agendamiento', f'{W}&id_medico={otro}')
        chk('filtro por otro médico: 0 citas y tasas en 0', kpi(d3, 'agendadas') == 0 and kpi(d3, 'tasa_asistencia') == 0)
    # una cita anulada (cabecera Inactivo) no cuenta
    cu.execute("SELECT count(*) FROM cita_detalle d JOIN cita_cabecera c USING (id_cita_cabecera) WHERE c.estado = 'Inactivo'")
    r, d4 = api('agendamiento', AMPLIA)
    esperado = uno("""SELECT count(*) FROM cita_detalle d JOIN cita_cabecera cc USING (id_cita_cabecera) WHERE cc.estado = 'Activo' AND d.fecha_cita BETWEEN '2020-01-01' AND '2026-12-31'""")
    chk('ventana amplia: agendadas = SQL independiente (cabeceras inactivas excluidas)', kpi(d4, 'agendadas') == esperado)
    esp_aus = uno("""SELECT count(*) FROM cita_detalle d JOIN cita_cabecera cc USING (id_cita_cabecera)
                      WHERE cc.estado = 'Activo' AND d.fecha_cita BETWEEN '2020-01-01' AND '2026-12-31'
                        AND (d.id_estado_cita = 5 OR (d.id_estado_cita = 2 AND d.fecha_cita < CURRENT_DATE
                             AND NOT EXISTS (SELECT 1 FROM consultas_cab k WHERE k.id_cita = d.id_cita_detalle AND k.activo)))""")
    chk('ventana amplia: ausentes = SQL independiente', kpi(d4, 'ausentes') == esp_aus)

    # =========================================================================================================================
    # 6. CONSULTORIO (sin pacientes ni contenido clínico)
    # =========================================================================================================================
    r, d = api('consultorio')
    chk('Consultorio: 200 con kpis, 2 gráficos y 5 tablas', r.status_code == 200 and d['success'] and len(d['data']['graficos']) == 2
        and [t['id'] for t in d['data']['tablas']] == ['diagnosticos', 'tratamientos', 'procedimientos', 'insumos', 'por_medico'])
    chk('consultas 1, pacientes distintos 1, nuevos 1 (su primera consulta cae en el período)', (kpi(d, 'consultas'), kpi(d, 'pacientes'), kpi(d, 'nuevos')) == (1, 1, 1))
    chk('promedio diario = 1 / 31 días = 0,03', kpi(d, 'promedio') == 0.03)
    pmed = tabla(d, 'por_medico')['filas']
    chk('consultas por médico: una fila con 1', len(pmed) == 1 and pmed[0]['cantidad'] == 1)
    r, dam = api('consultorio', 'desde=2026-01-01&hasta=2026-12-31')
    esp = uno("""SELECT count(DISTINCT id_paciente) FROM (SELECT id_paciente, MIN(fecha_cita) OVER (PARTITION BY id_paciente) pf, fecha_cita
                   FROM consultas_cab WHERE activo AND estado <> 'cancelada') x WHERE fecha_cita BETWEEN '2026-01-01' AND '2026-12-31' AND pf BETWEEN '2026-01-01' AND '2026-12-31'""")
    chk('el paciente con consulta en 2001 NO es nuevo en 2026 (nuevos = SQL independiente)', kpi(dam, 'nuevos') == esp)
    esp_cons = uno("SELECT count(*) FROM consultas_cab WHERE activo AND estado <> 'cancelada' AND fecha_cita BETWEEN '2026-01-01' AND '2026-12-31'")
    chk('consultas 2026 = SQL independiente', kpi(dam, 'consultas') == esp_cons)
    esp_tr = dict(uno_fila for uno_fila in (cu.execute("SELECT estado, count(*) FROM tratamientos WHERE fecha_tratamiento BETWEEN '2020-01-01' AND '2026-12-31' GROUP BY 1") or cu.fetchall()))
    r, dw = api('consultorio', AMPLIA)
    chk('tratamientos por estado = SQL independiente', {f['nombre']: f['cantidad'] for f in tabla(dw, 'tratamientos')['filas']} == esp_tr)
    cu.execute("SELECT nombre, apellido FROM paciente WHERE activo")
    nombres = [f'{n} {a}'.lower() for n, a in cu.fetchall()] + [a.lower() for n, a in [] ]
    cu.execute("SELECT nombre, apellido, cedula_entidad FROM paciente")
    pers = cu.fetchall()
    cuerpo = r.get_data(as_text=True).lower()
    chk('el informe de Consultorio no contiene nombres, apellidos ni CI de pacientes', not any((n.lower() in cuerpo and len(n) > 3) or (a.lower() in cuerpo and len(a) > 3) or (ci and ci in cuerpo) for n, a, ci in pers
        if n.lower() not in ('de', 'la') and a.lower() not in ('de', 'la') and not any(w.lower() in ('consultas', 'pacientes') for w in (n, a))))
    cu.execute("SELECT id_medico FROM medico WHERE activo ORDER BY 1")
    ms = [m[0] for m in cu.fetchall()]
    r, dmf = api('consultorio', f'{AMPLIA}&id_medico={ms[0]}')
    esp_m = uno("SELECT count(*) FROM consultas_cab WHERE activo AND estado <> 'cancelada' AND id_medico = %s AND fecha_cita BETWEEN '2020-01-01' AND '2026-12-31'", (ms[0],))
    chk('filtro por médico: consultas = SQL independiente', kpi(dmf, 'consultas') == esp_m)

    # =========================================================================================================================
    # 7. FACTURACIÓN
    # =========================================================================================================================
    r, d = api('facturacion', AMPLIA)
    chk('Facturación: 200 con 5 kpis, 2 gráficos y 4 tablas', r.status_code == 200 and d['success'] and len(d['data']['kpis']) == 5 and len(d['data']['graficos']) == 2
        and [t['id'] for t in d['data']['tablas']] == ['condicion', 'cobros', 'cuentas', 'detalle'])
    print('   KPIs:', {k['clave']: k['valor'] for k in d['data']['kpis']})
    esp = uno("""SELECT COALESCE((SELECT SUM(total) FROM venta_cab WHERE estado = 'EMITIDA' AND fecha::date BETWEEN '2020-01-01' AND '2026-12-31'), 0)
                      - COALESCE((SELECT SUM(total) FROM nota_credito_debito_cab WHERE estado = 'EMITIDA' AND tipo_nota = 'CREDITO' AND fecha::date BETWEEN '2020-01-01' AND '2026-12-31'), 0)
                      + COALESCE((SELECT SUM(total) FROM nota_credito_debito_cab WHERE estado = 'EMITIDA' AND tipo_nota = 'DEBITO' AND fecha::date BETWEEN '2020-01-01' AND '2026-12-31'), 0)""")
    chk('total neto = facturas - NC + ND (SQL independiente sobre las tablas)', kpi(d, 'total') == esp)
    esp10 = uno("SELECT COALESCE(SUM(total_gravada_10), 0) FROM venta_cab WHERE estado = 'EMITIDA' AND fecha::date BETWEEN '2020-01-01' AND '2026-12-31'")
    chk('gravado 10 % = SQL independiente (sin notas en la base de prueba)', kpi(d, 'gravado_10') == esp10 or uno("SELECT count(*) FROM nota_credito_debito_cab WHERE estado = 'EMITIDA'") > 0)
    chk('total = gravado 10 % + gravado 5 % + exentas', kpi(d, 'total') == kpi(d, 'gravado_10') + kpi(d, 'gravado_5') + kpi(d, 'exentas'))
    chk('IVA = IVA 5 % + IVA 10 % de los comprobantes', kpi(d, 'iva') == uno("SELECT COALESCE(SUM(iva_5 + iva_10), 0) FROM v_inf_comprobantes WHERE estado = 'EMITIDA' AND fecha BETWEEN '2020-01-01' AND '2026-12-31'"))
    cond = {f['condicion']: f for f in tabla(d, 'condicion')['filas']}
    chk('contado + crédito = total', sum(f['total'] for f in cond.values()) == kpi(d, 'total'))
    r, dco = api('facturacion', f'{AMPLIA}&condicion=CONTADO')
    r, dcr = api('facturacion', f'{AMPLIA}&condicion=CREDITO')
    chk('filtro por condición: contado y crédito suman el total y cada uno coincide con su fila', kpi(dco, 'total') + kpi(dcr, 'total') == kpi(d, 'total')
        and kpi(dco, 'total') == cond.get('Contado', {'total': 0})['total'] and kpi(dcr, 'total') == cond.get('Credito', {'total': 0})['total']
        and all(f['condicion'] == 'Contado' for f in tabla(dco, 'detalle')['filas']))
    esp_cobros = uno("SELECT COALESCE(SUM(d.monto), 0) FROM cobro_det d JOIN cobro_cab c USING (id_cobro) WHERE c.estado = 'EMITIDO' AND c.fecha::date BETWEEN '2020-01-01' AND '2026-12-31'")
    chk('cobros por forma de cobro suman lo cobrado (SQL independiente)', sum(f['monto'] for f in tabla(d, 'cobros')['filas']) == esp_cobros)
    esp_saldo = uno("""SELECT COALESCE(SUM(cc.saldo), 0) FROM cuenta_cobrar cc JOIN venta_cab v USING (id_venta) WHERE cc.estado = 'PENDIENTE' AND cc.saldo > 0
                         AND v.fecha::date BETWEEN '2020-01-01' AND '2026-12-31'""")
    chk('cuentas a cobrar: suma de saldos = SQL independiente', sum(f['saldo'] for f in tabla(d, 'cuentas')['filas']) == esp_saldo and all(f['saldo'] > 0 for f in tabla(d, 'cuentas')['filas']))
    det = tabla(d, 'detalle')['filas']
    chk('detalle de comprobantes: solo EMITIDAS y suma = total', sum(f['total'] for f in det) == kpi(d, 'total')
        and len(det) == uno("SELECT count(*) FROM v_inf_comprobantes WHERE estado = 'EMITIDA' AND fecha BETWEEN '2020-01-01' AND '2026-12-31'"))
    ventas_dia = d['data']['graficos'][0]
    chk('ventas por día suman el total', sum(ventas_dia['series'][0]['datos']) == kpi(d, 'total'))
    # signo de las notas: se prueba la vista dentro de una transacción que se revierte
    cv = Conexion().getConexion(); cv.autocommit = False; cvu = cv.cursor()
    try:
        cvu.execute("SET LOCAL session_replication_role = replica")     # sin triggers: solo se prueba la vista; todo se revierte
        cvu.execute("""SELECT cc.id_cuenta, v.id_timbrado, v.id_usuario, v.id_apertura, v.condicion FROM cuenta_cobrar cc JOIN venta_cab v USING (id_venta)
                        WHERE v.estado = 'EMITIDA' ORDER BY cc.id_cuenta LIMIT 1""")
        cuenta, timb, usr, aper, condv = cvu.fetchone()
        for tipo, signo, num in (('CREDITO', -1, 'ZZ-NC'), ('DEBITO', 1, 'ZZ-ND')):
            cvu.execute("""INSERT INTO nota_credito_debito_cab (id_cuenta, id_timbrado, id_usuario, tipo_nota, numero_nota, fecha, motivo, total_exenta, total_gravada_5,
                                                                total_gravada_10, iva_5, iva_10, total, estado, id_apertura, monto_reintegro)
                           VALUES (%s, %s, %s, %s, %s, '2001-03-15', 'ZZ', 0, 0, 1100, 0, 100, 1100, 'EMITIDA', %s, 0)""", (cuenta, timb, usr, tipo, num, aper))
            cvu.execute("SELECT tipo, total, gravada_10, iva_10, condicion FROM v_inf_comprobantes WHERE numero = %s", (num,))
            fila = cvu.fetchone()
            chk(f'vista: la nota de {"crédito resta" if signo < 0 else "débito suma"} (total, gravada 10 % e IVA con signo) y toma la condición de la venta',
                fila is not None and (fila[1], fila[2], fila[3]) == (1100 * signo, 1100 * signo, 100 * signo) and fila[4] == condv)
    finally:
        cv.rollback(); cvu.close(); cv.close()

    # =========================================================================================================================
    # 8. AUDITORÍA (accesos)
    # =========================================================================================================================
    r, d = api('auditoria', AMPLIA)
    chk('Auditoría: 200 con 3 kpis, 1 gráfico y 2 tablas', r.status_code == 200 and d['success'] and [k['clave'] for k in d['data']['kpis']] == ['exitosos', 'fallidos', 'bloqueos']
        and len(d['data']['graficos']) == 1 and [t['id'] for t in d['data']['tablas']] == ['ips', 'usuarios'])
    print('   KPIs:', {k['clave']: k['valor'] for k in d['data']['kpis']})
    chk('accesos exitosos = LOGIN EXITOSO (SQL independiente)', kpi(d, 'exitosos') == uno("SELECT count(*) FROM auditoria_acceso WHERE accion = 'LOGIN' AND resultado = 'EXITOSO' AND fecha_hora::date BETWEEN '2020-01-01' AND '2026-12-31'"))
    chk('accesos fallidos = LOGIN / 2FA no exitosos (SQL independiente)', kpi(d, 'fallidos') == uno("""SELECT count(*) FROM auditoria_acceso WHERE accion IN ('LOGIN', 'VERIFICACION_2FA')
                                                                                                          AND resultado <> 'EXITOSO' AND fecha_hora::date BETWEEN '2020-01-01' AND '2026-12-31'"""))
    chk('bloqueos = acciones BLOQUEO (SQL independiente)', kpi(d, 'bloqueos') == uno("SELECT count(*) FROM auditoria_acceso WHERE accion = 'BLOQUEO' AND fecha_hora::date BETWEEN '2020-01-01' AND '2026-12-31'"))
    chk('bloqueos por IP suman los bloqueos', sum(f['bloqueos'] for f in tabla(d, 'ips')['filas']) == kpi(d, 'bloqueos'))
    uf = tabla(d, 'usuarios')['filas']
    chk('usuarios con más intentos fallidos: ordenados de mayor a menor', [f['fallidos'] for f in uf] == sorted([f['fallidos'] for f in uf], reverse=True) and len(uf) >= 1)
    chk('un período sin accesos devuelve ceros y tablas vacías', (lambda dd: kpi(dd, 'exitosos') == 0 and kpi(dd, 'fallidos') == 0 and tabla(dd, 'ips')['filas'] == [])(api('auditoria', W)[1]))

    # =========================================================================================================================
    # 9. RESUMEN
    # =========================================================================================================================
    r, d = api('resumen', AMPLIA)
    chk('Resumen: 200 con 5 kpis, gráfico de ingresos y sin tablas', r.status_code == 200 and [k['clave'] for k in d['data']['kpis']] == ['total_facturado', 'saldo_por_cobrar', 'citas', 'tasa_asistencia', 'consultas']
        and d['data']['tablas'] == [] and d['data']['graficos'][0]['id'] == 'ingresos_dia')
    r, df = api('facturacion', AMPLIA)
    chk('Resumen: total facturado = total de Facturación', kpi(d, 'total_facturado') == kpi(df, 'total'))
    r, da = api('agendamiento', AMPLIA)
    chk('Resumen: citas y tasa de asistencia = las de Agendamiento', kpi(d, 'citas') == kpi(da, 'agendadas') and kpi(d, 'tasa_asistencia') == kpi(da, 'tasa_asistencia'))
    r, dc = api('consultorio', AMPLIA)
    chk('Resumen: consultas = las de Consultorio', kpi(d, 'consultas') == kpi(dc, 'consultas'))
    chk('Resumen: ingresos por día suman lo cobrado', sum(d['data']['graficos'][0]['series'][0]['datos']) == esp_cobros)
    chk('Resumen: saldo por cobrar = suma de saldos de las cuentas PENDIENTES', kpi(d, 'saldo_por_cobrar') == esp_saldo)

    # =========================================================================================================================
    # 10. AUDITORÍA DE LAS CONSULTAS
    # =========================================================================================================================
    cu.execute("""SELECT detalle FROM auditoria_cambios WHERE id_auditoria_cambio > %s AND entidad = 'informes' AND accion = 'CONSULTA' AND id_usuario = %s""", (BASE, ADMIN[0]))
    det_aud = [x[0] for x in cu.fetchall()]
    chk('cada pestaña consultada quedó auditada con informe y filtros', all(any(x.startswith(f'Informe: {n};') for x in det_aud) for n in ('Resumen', 'Agendamiento', 'Consultorio', 'Facturación', 'Auditoría')))
    chk('el médico filtrado se audita con su nombre (no con el ID)', any('Médico: ' in x and 'N.º' not in x for x in det_aud))
    chk('la condición filtrada se audita', any('Condición: Contado' in x for x in det_aud))
    chk('las consultas con error de filtros NO se auditan como informe', not any('2026-10-10' in x for x in det_aud))
    # una consulta de Informes (solo lectura) de madrugada no debe disparar la alerta FUERA_DE_HORARIO
    cu.execute("SELECT count(*) FROM v_auditoria_anomalias WHERE tipo = 'FUERA_DE_HORARIO'"); fuera_antes = cu.fetchone()[0]
    cu.execute("""INSERT INTO auditoria_cambios (fecha_hora, id_usuario, entidad, id_registro, accion, detalle, ip, modulo)
                  VALUES (CURRENT_DATE - 1 + TIME '03:00', %s, 'informes', 0, 'CONSULTA', 'ZZ madrugada', '127.0.0.1', 'Módulo Informes')""", (ADMIN[0],))
    cu.execute("SELECT count(*) FROM v_auditoria_anomalias WHERE tipo = 'FUERA_DE_HORARIO'")
    chk('una consulta de Informes fuera de horario NO genera la alerta FUERA_DE_HORARIO', cu.fetchone()[0] == fuera_antes)

    # =========================================================================================================================
    # 11. ETAPA 3: EXPORTAR PDF Y EXCEL
    # =========================================================================================================================
    from openpyxl import load_workbook
    from tests.facturacion._pdf import paginas_pdf, texto_pdf
    from app.Services import informes_export_service as ex
    cu.execute("SELECT COALESCE(MAX(id_auditoria_cambio), 0) FROM auditoria_cambios"); BASE_EXP = cu.fetchone()[0]
    MIME_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    TABLAS = {'agendamiento': ['Citas por médico', 'Pacientes ausentes', 'Ausencias reiteradas (2 o más)', 'Citas canceladas'],
              'consultorio': ['Diagnósticos más frecuentes', 'Tratamientos por estado', 'Procedimientos más usados', 'Insumos más usados (unidades)', 'Consultas por médico'],
              'facturacion': ['Contado vs. crédito', 'Cobros por forma de cobro', 'Cuentas a cobrar con saldo', 'Detalle de comprobantes (notas de crédito en negativo)'],
              'auditoria': ['Bloqueos e intentos fallidos por IP', 'Usuarios con más intentos fallidos']}
    FILTRO = {'agendamiento': W, 'consultorio': W, 'facturacion': AMPLIA, 'auditoria': AMPLIA}
    detalle_de = lambda libro: libro[next(h for h in libro.sheetnames if h.startswith('Detalle'))]      # el nombre de hoja se corta a 28 caracteres
    exportados = {}
    for pest, titulos in TABLAS.items():
        rp = c.get(f'/api/informes/{pest}/exportar?formato=pdf&{FILTRO[pest]}')
        chk(f'{pest}: PDF 200, application/pdf, empieza con %PDF y se ve en el navegador', rp.status_code == 200 and rp.mimetype == 'application/pdf'
            and rp.data[:5] == b'%PDF-' and 'inline' in rp.headers['Content-Disposition'] and f'informe_{pest}_' in rp.headers['Content-Disposition'])
        txt = texto_pdf(rp.data)
        chk(f'{pest}: el PDF lleva título, filtros e indicadores', f'INFORME DE {pest.upper()[:5]}' in txt.upper() and 'Filtros aplicados' in txt and 'Indicadores' in txt)
        rx = c.get(f'/api/informes/{pest}/exportar?formato=xlsx&{FILTRO[pest]}')
        chk(f'{pest}: Excel 200 con el tipo MIME de xlsx y como descarga', rx.status_code == 200 and rx.mimetype == MIME_XLSX and 'attachment' in rx.headers['Content-Disposition']
            and rx.headers['Content-Disposition'].endswith('.xlsx'))
        libro = load_workbook(io.BytesIO(rx.data))
        exportados[pest] = libro
        chk(f'{pest}: el Excel tiene la hoja Indicadores y una hoja por tabla', libro.sheetnames[0] == 'Indicadores' and len(libro.sheetnames) == 1 + len(titulos)
            and all(any(h.startswith(t[:20]) for h in libro.sheetnames[1:]) for t in titulos))
    ind = exportados['agendamiento']['Indicadores']
    filas_ind = {f[0].value: f[1].value for f in ind.iter_rows() if f[0].value}
    chk('Excel de Agendamiento: indicadores numéricos reales (agendadas 6, ausentes 2, asistencia 33,3)',
        filas_ind.get('Agendadas') == 6 and filas_ind.get('Ausentes') == 2 and filas_ind.get('Tasa de asistencia') == 33.3 and isinstance(filas_ind.get('Agendadas'), (int, float)))
    chk('Excel de Agendamiento: lleva el período filtrado y quién lo emitió', any(r[0].value == 'Filtro' and r[1].value == 'Período' and r[2].value == '01/03/2001 al 31/03/2001' for r in ind.iter_rows())
        and any(r[0].value == 'Emitido' and r[3].value == 'USUARIO PRUEBA' for r in ind.iter_rows()))
    h_aus = exportados['agendamiento']['Pacientes ausentes']
    datos_aus = [[x.value for x in f] for f in h_aus.iter_rows(min_row=4)]
    chk('Excel de Agendamiento: hoja Pacientes ausentes con 2 filas y ausencias previas numéricas', len(datos_aus) == 2 and sorted(f[-1] for f in datos_aus) == [0, 1])
    h_med = exportados['agendamiento']['Citas por médico']
    fila_med = [x.value for x in next(h_med.iter_rows(min_row=4))]
    chk('Excel: encabezado en negrita, panel fijo y porcentaje como número', h_med['A3'].font.bold and h_med.freeze_panes == 'A4' and fila_med[-1] == 33.3 and '%' in h_med.cell(row=4, column=6).number_format)
    h_fac = detalle_de(exportados['facturacion'])
    ced = [x.value for x in next(h_fac.iter_rows(min_row=4))]
    chk('Excel de Facturación: importes como números con formato de miles', isinstance(ced[-1], (int, float)) and h_fac.cell(row=4, column=11).number_format == '#,##0')
    cu.execute("SELECT nombre, apellido FROM paciente")
    nombres_pac = [f'{n} {a}'.lower() for n, a in cu.fetchall()]
    texto_xlsx = ' '.join(str(x.value).lower() for h in exportados['consultorio'].worksheets for f in h.iter_rows() for x in f if x.value is not None)
    chk('el Excel de Consultorio no contiene nombres de pacientes', not any(n in texto_xlsx for n in nombres_pac))
    pdf_c = c.get(f'/api/informes/consultorio/exportar?formato=pdf&{AMPLIA}').data
    chk('el PDF de Consultorio no contiene nombres de pacientes', not any(n in texto_pdf(pdf_c).lower() for n in nombres_pac))
    # filtros en el archivo
    cu.execute("SELECT nombre, apellido FROM medico WHERE id_medico = %s", (id_medico,))
    nm = ' '.join(cu.fetchone())
    rp = c.get(f'/api/informes/agendamiento/exportar?formato=pdf&{W}&id_medico={id_medico}')
    chk('el PDF indica el médico filtrado por nombre', nm.split()[0] in texto_pdf(rp.data))
    rf = c.get(f'/api/informes/facturacion/exportar?formato=xlsx&{AMPLIA}&condicion=CONTADO')
    chk('el Excel de Facturación filtrado por condición solo trae comprobantes de contado', all(f[5].value == 'Contado' for f in detalle_de(load_workbook(io.BytesIO(rf.data))).iter_rows(min_row=4)))
    # validaciones
    chk('Resumen no se exporta (400)', c.get(f'/api/informes/resumen/exportar?formato=pdf&{W}').status_code == 400)
    chk('formato inválido -> 400', c.get(f'/api/informes/agendamiento/exportar?formato=csv&{W}').status_code == 400 and c.get(f'/api/informes/agendamiento/exportar?{W}').status_code == 400)
    chk('filtros inválidos al exportar -> 400', c.get('/api/informes/agendamiento/exportar?formato=pdf&desde=2026-10-10&hasta=2026-10-01').status_code == 400
        and c.get(f'/api/informes/facturacion/exportar?formato=pdf&{W}&condicion=x').status_code == 400)
    chk('pestaña inexistente al exportar -> 404', c.get(f'/api/informes/otra/exportar?formato=pdf&{W}').status_code == 404)
    como(CAJERO)
    chk('Cajero no puede exportar (403)', c.get(f'/api/informes/agendamiento/exportar?formato=pdf&{W}').status_code == 403 and c.get(f'/api/informes/agendamiento/exportar?formato=xlsx&{W}').status_code == 403)
    como(None)
    chk('sin sesión no se puede exportar (401)', c.get(f'/api/informes/agendamiento/exportar?formato=pdf&{W}').status_code == 401)
    como(ADMIN)
    # tablas largas: paginación del PDF y filas completas del Excel
    grande = {'kpis': [{'clave': 'n', 'titulo': 'Filas', 'valor': 400, 'formato': 'num'}],
              'tablas': [{'id': 't', 'titulo': 'Tabla larga', 'columnas': [{'campo': 'a', 'titulo': 'Texto', 'tipo': 'texto'}, {'campo': 'b', 'titulo': 'Monto', 'tipo': 'gs'}],
                          'filas': [{'a': f'Fila {i}', 'b': i * 1000} for i in range(400)] + [{'a': '=SUMA(1+1)', 'b': 1}]}]}
    armado = ex.calcular_datos_exportacion('Prueba', {'razon_social': 'EMPRESA ZZ', 'ruc': '1-1', 'direccion': 'X', 'telefono': '1'}, [('Período', 'p')], grande, 'ZZ')
    pdf_g = ex.generarInformePdf(armado)
    chk('PDF largo: varias páginas con la sección repetida como continuación', paginas_pdf(pdf_g) > 3 and 'continuaci' in texto_pdf(pdf_g))
    hg = load_workbook(io.BytesIO(ex.generarInformeExcel(armado)))['Tabla larga']
    chk('Excel largo: todas las filas (401 + 3 de encabezado)', hg.max_row == 3 + 401)
    chk('Excel: un texto que empieza con "=" no queda como fórmula', hg.cell(row=hg.max_row, column=1).data_type != 'f' and not str(hg.cell(row=hg.max_row, column=1).value).startswith('='))
    chk('formato de números del PDF: miles con punto, porcentaje con coma', ex.formatear(1234567, 'gs') == 'Gs. 1.234.567' and ex.formatear(33.3, 'pct') == '33,3 %' and ex.formatear(1234, 'num') == '1.234')
    chk('aviso de recorte cuando una tabla alcanza el tope', ex.aviso_de_recorte([{'titulo': 'T', 'filas': [0] * ex.TOPE_FILAS}]) != '' and ex.aviso_de_recorte([{'titulo': 'T', 'filas': [0]}]) == '')
    # auditoría de las exportaciones
    cu.execute("SELECT detalle, modulo, id_usuario, ip FROM auditoria_cambios WHERE id_auditoria_cambio > %s AND entidad = 'informes' AND accion = 'EXPORTACION' ORDER BY id_auditoria_cambio", (BASE_EXP,))
    aud = cu.fetchall()
    print('   exportaciones auditadas:', len(aud))
    chk('cada exportación exitosa quedó auditada (8 + PDF Consultorio + PDF por médico + Excel contado = 11; las rechazadas no)', len(aud) == 11 and all(a[1] == 'Módulo Informes' and a[2] == ADMIN[0] and a[3] == '127.0.0.1' for a in aud))
    chk('la auditoría lleva formato, informe, filtros y cantidad de filas', any(a[0].startswith('PDF; Informe: Agendamiento; Período: 01/03/2001 al 31/03/2001; ') and a[0].endswith(' filas') for a in aud)
        and any(a[0].startswith('XLSX; Informe: Facturación; Período: 01/01/2020 al 31/12/2026; Condición: Contado; ') for a in aud)
        and any(f'Médico: {nm}' in a[0] for a in aud))

finally:
    # Limpieza: SOLO lo creado por este test, por ID.
    if consultas_zz:
        cu.execute("DELETE FROM consultas_cab WHERE id_consulta_cab = ANY(%s)", (consultas_zz,))
    if citas_zz:
        cu.execute("DELETE FROM cita_detalle WHERE id_cita_detalle = ANY(%s)", (citas_zz,))
    cu.execute("SELECT id_auditoria_cambio FROM auditoria_cambios WHERE id_auditoria_cambio > %s AND entidad = 'informes' AND id_usuario = %s AND ip = '127.0.0.1'", (BASE, ADMIN[0]))
    ids = [x[0] for x in cu.fetchall()]
    cu.execute("SET session_replication_role = replica")
    cu.execute("DELETE FROM auditoria_cambios WHERE id_auditoria_cambio = ANY(%s)", (ids,))
    borradas = cu.rowcount
    cu.execute("RESET session_replication_role")
    print(f'(limpieza: {len(citas_zz)} citas, {len(consultas_zz)} consulta y {borradas} filas de auditoría de prueba borradas por ID)')
    cu.close(); cx.close()

print('FALLOS:', 'ninguno' if not fallos else fallos)
