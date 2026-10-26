"""Formato de montos y cantidades en todo el sistema: guaraníes sin decimales con punto de miles (100000 -> 100.000), un solo helper para mostrar
(app/Services/formato.py + filtros Jinja monto / gs / cantidad, espejo de formatearMonto / limpiarMonto en facturacion-validadores.js), validación
estricta del servidor (entero >= 0, o > 0 donde corresponde; llega SIN puntos), pantallas con campos data-monto / data-cantidad (sin maxlength, que contaría
los puntos), PDFs e Informes (en Excel como número, con formato de miles, no como texto).
El comportamiento EN VIVO del campo (cursor, letras, tope de dígitos) y el envío del entero están en browser_montos.py; la comparación JS <-> Python en
comparar_js_py.py (casos 'montoFmt', 'montoLimp', 'gsFmt').
Datos de prueba con prefijo ZZ, borrados POR ID; la auditoría que genera el test se borra por ID."""
import io, pathlib, re, sys
RAIZ = pathlib.Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RAIZ))
from app import app
from tests._conexion_pruebas import Conexion
from tests.facturacion._pdf import texto_pdf
from app.Services import formato as fm
from app.Services import factura_pdf_service, informes_export_service as ex
from app.rutas.modulo_facturacion import utils as u

fallos = []
c = app.test_client()


def chk(nombre, cond):
    print(('OK  ' if cond else 'FAIL'), nombre)
    if not cond:
        fallos.append(nombre)


print('== Helper único de formato (Python)')
CASOS = {100000: '100.000', 0: '0', 999: '999', 1000: '1.000', 15000: '15.000', 999999999999: '999.999.999.999', -1500: '-1.500', 1234567: '1.234.567',
         '15000': '15.000', '100.000': '100.000', 1234.6: '1.235', 1234.4: '1.234', None: '', '': '', '   ': ''}
from decimal import Decimal
CASOS[Decimal('2500000.00')] = '2.500.000'
for entrada, esperado in CASOS.items():
    chk(f'formatear_monto({entrada!r}) -> {esperado!r}', fm.formatear_monto(entrada) == esperado)
chk('formatear_monto de un texto no numérico lo devuelve como está', fm.formatear_monto('abc') == 'abc')
chk('formatear_gs: prefijo "Gs." y vacío', fm.formatear_gs(1500000) == 'Gs. 1.500.000' and fm.formatear_gs(0) == 'Gs. 0' and fm.formatear_gs(None) == '')
chk('formatear_cantidad = formatear_monto', fm.formatear_cantidad(1500) == '1.500' and fm.formatear_cantidad(5) == '5')
chk('limpiar_monto: solo dígitos y sin ceros a la izquierda', fm.limpiar_monto('100.000') == '100000' and fm.limpiar_monto('Gs. 1.500') == '1500' and fm.limpiar_monto('007') == '7'
    and fm.limpiar_monto('000') == '0' and fm.limpiar_monto('') == '' and fm.limpiar_monto(None) == '' and fm.limpiar_monto('abc') == '')
chk('ida y vuelta: limpiar(formatear(n)) == str(n)', all(fm.limpiar_monto(fm.formatear_monto(n)) == str(n) for n in (0, 5, 999, 1000, 100000, 123456789, 999999999999)))
chk('no queda ningún `def _gs` duplicado en la aplicación (todos usan formato.py)', not any('def _gs(' in p.read_text(encoding='utf-8') for p in (RAIZ / 'app').rglob('*.py')))
chk('los 7 módulos que tenían su propio _gs usan el helper único', all('from app.Services.formato import formatear_gs as _gs' in (RAIZ / r).read_text(encoding='utf-8') for r in (
    'app/rutas/modulo_facturacion/apertura_caja/apertura_caja_api.py', 'app/rutas/modulo_facturacion/arqueo/arqueo_api.py', 'app/rutas/modulo_facturacion/cobro/cobro_api.py',
    'app/rutas/modulo_facturacion/nota_credito_debito/nota_api.py', 'app/rutas/modulo_facturacion/pedido_cliente/pedido_cliente_api.py',
    'app/rutas/modulo_facturacion/venta/venta_api.py', 'app/Services/pedido_pdf_service.py')))

print('\n== Filtros Jinja')
f = app.jinja_env.filters
chk('|monto, |gs y |cantidad', f['monto'](100000) == '100.000' and f['gs'](100000) == 'Gs. 100.000' and f['cantidad'](1500) == '1.500' and f['monto'](None) == '')
chk('en una plantilla: {{ n|gs }}', app.jinja_env.from_string('{{ 2500000|gs }} / {{ 1234|cantidad }} / {{ x|monto }}').render(x=None) == 'Gs. 2.500.000 / 1.234 / ')

print('\n== El servidor valida entero >= 0 (el valor llega sin puntos)')
for v in (0, 1, 100000, '100000', 999999999999):
    chk(f'validar_monto({v!r}) acepta', u.validar_monto(v)[1] is None)
for v in ('100.000', '1,5', '1.5', 'abc', '-1', -5, '', None, 1000000000000, '1e3', '0x10', ' 1 000'):
    chk(f'validar_monto({v!r}) rechaza', u.validar_monto(v)[1] is not None)
chk('validar_precio: 0 se rechaza (mayor a 0) y 1 se acepta', u.validar_precio(0)[1] is not None and u.validar_precio('0')[1] is not None and u.validar_precio(1)[1] is None)
chk('validar_precio rechaza "100.000" (con puntos) y decimales', u.validar_precio('100.000')[1] is not None and u.validar_precio('150.5')[1] is not None and u.validar_precio('150,5')[1] is not None)
chk('validar_monto_positivo: 0 se rechaza', u.validar_monto_positivo(0)[1] is not None and u.validar_monto_positivo('1500')[1] is None and u.validar_monto_positivo('1.500')[1] is not None)
chk('validar_cantidad: 0 se acepta (stock 0), negativa y con puntos no', u.validar_cantidad('0')[1] is None and u.validar_cantidad('-1')[1] is not None and u.validar_cantidad('1.500')[1] is not None)
chk('validar_cantidad_positiva: 0 se rechaza', u.validar_cantidad_positiva('0')[1] is not None and u.validar_cantidad_positiva('3')[1] is None)
chk('los mensajes con topes usan puntos de miles', '999.999.999.999' in u.validar_monto(10**12)[1] and '2.147.483.647' in u.validar_cantidad(2**31)[1])

cx = Conexion().getConexion(); cx.autocommit = True; cu = cx.cursor()
cu.execute("SELECT COALESCE(MAX(id_auditoria_cambio), 0) FROM auditoria_cambios"); BASE = cu.fetchone()[0]
creados = {'tipos_tratamiento': [], 'items': []}


def api(metodo, url, cuerpo):
    r = getattr(c, metodo)(url, json=cuerpo)
    return r, (r.get_json(silent=True) or {})


try:
    with c.session_transaction() as s:
        s['id_usuario'] = 1; s['id_rol'] = 1; s['nombre_rol'] = 'Administrador'

    print('\n== APIs (precio de Tipo Tratamiento e Ítems)')
    r, d = api('post', '/api/v1/tipos-tratamiento', {'descripcion': 'zz montos uno', 'precio': '1500000', 'iva': '10'})
    chk('Tipo Tratamiento: precio "1500000" -> 201 y se guarda el entero', r.status_code == 201)
    if r.status_code == 201:
        i = d['data']['id_tipo_tratamiento']; creados['tipos_tratamiento'].append(i)
        cu.execute("SELECT precio FROM tipos_tratamiento WHERE id_tipo_tratamiento = %s", (i,))
        chk('Tipo Tratamiento: en la base 1500000', int(cu.fetchone()[0]) == 1500000)
    for malo, texto in (('1.500.000', 'con puntos'), ('1500,50', 'con coma'), ('1500.5', 'decimal'), ('abc', 'letras'), ('0', 'cero'), ('-5', 'negativo'), ('', 'vacío')):
        r, d = api('post', '/api/v1/tipos-tratamiento', {'descripcion': 'zz montos dos', 'precio': malo, 'iva': '10'})
        chk(f'Tipo Tratamiento: precio {malo!r} ({texto}) -> 400 en el campo precio', r.status_code == 400 and d.get('campo') == 'precio')
    cu.execute("SELECT id_tipo_item FROM tipo_items WHERE activo AND controla_stock ORDER BY 1 LIMIT 1"); tipo_item = cu.fetchone()
    cu.execute("SELECT id_deposito FROM deposito WHERE activo ORDER BY 1 LIMIT 1"); deposito = cu.fetchone()
    if tipo_item and deposito:
        base = {'id_tipo_item': tipo_item[0], 'id_deposito': deposito[0], 'descripcion': 'zz item montos', 'precio_unitario': '2500000', 'iva': '10', 'cantidad_disponible': '1500'}
        r, d = api('post', '/api/v1/items', base)
        chk('Ítem: precio "2500000" y stock "1500" -> 201', r.status_code == 201)
        if r.status_code == 201:
            i = d['data']['id_item']; creados['items'].append(i)
            cu.execute("SELECT precio_unitario, cantidad_disponible FROM items WHERE id_item = %s", (i,))
            fila = cu.fetchone()
            chk('Ítem: en la base 2500000 y 1500 (enteros)', (int(fila[0]), int(fila[1])) == (2500000, 1500))
        for campo, malo in (('precio_unitario', '2.500.000'), ('precio_unitario', '2500,5'), ('precio_unitario', '0'), ('cantidad_disponible', '1.500'), ('cantidad_disponible', '-1'), ('cantidad_disponible', 'x')):
            r, d = api('post', '/api/v1/items', {**base, 'descripcion': 'zz item montos 2', campo: malo})
            chk(f'Ítem: {campo} {malo!r} -> 400 en ese campo', r.status_code == 400 and d.get('campo') == campo)
    else:
        chk('hay un tipo de ítem con stock y un depósito para probar Ítems', False)

    print('\n== Pantallas: campos data-monto / data-cantidad, sin maxlength, y el entero sale sin puntos')
    PAGINAS = [('/referenciales/tipo-tratamiento/tipo-tratamiento-index', 1, 0), ('/referenciales/item/item-index', 1, 1),
               ('/referenciales/apertura-caja/apertura-caja-index', 1, 0), ('/referenciales/arqueo/arqueo-index', 1, 0), ('/referenciales/cobro/cobro-index', 1, 0),
               ('/referenciales/nota-credito-debito/nota-index', 1, 1), ('/referenciales/pedido-cliente/pedido-cliente-index', 0, 1),
               ('/referenciales/remision/remision-index', 0, 1)]
    for url, n_monto, n_cant in PAGINAS:
        r = c.get(url)
        html = r.get_data(as_text=True)
        tags = re.findall(r'<input[^>]*\bdata-(monto|cantidad)\b[^>]*>', html)
        montos = [t for t in tags if t == 'monto']; cants = [t for t in tags if t == 'cantidad']
        sin_maxlength = all('maxlength' not in m.group(0) for m in re.finditer(r'<input[^>]*\bdata-(?:monto|cantidad)\b[^>]*>', html))
        chk(f'{url.split("/")[2]}: carga, {n_monto} campo(s) de monto y {n_cant} de cantidad, ninguno con maxlength',
            r.status_code == 200 and len(montos) == n_monto and len(cants) == n_cant and sin_maxlength and 'facturacion-validadores.js' in html)
        chk(f'{url.split("/")[2]}: formatea con formatearMonto/formatoGs y limpia con limpiarMonto / valorCampo',
            ('formatearMonto' in html or 'formatoGs' in html) and ('limpiarMonto' in html or 'valorCampo' in html))
    js = c.get('/static/js/facturacion-validadores.js').get_data(as_text=True)
    chk('el JS expone formatearMonto, limpiarMonto, valorCampo y el formateo en vivo (captura, sin saltos de cursor)',
        all(x in js for x in ('formatearMonto', 'limpiarMonto', 'valorCampo', 'formatearEnVivo', 'setSelectionRange', "addEventListener('input'")))
    chk('no hay `toLocaleString` ni `toFixed` para montos en los scripts de la aplicación', 'toLocaleString' not in js and 'toFixed' not in js
        and 'toLocaleString' not in c.get('/static/js/informes.js').get_data(as_text=True).split('function formatear(')[0])
    html = c.get('/referenciales/tipo-tratamiento/tipo-tratamiento-index').get_data(as_text=True)
    chk('Tipo Tratamiento: precio con data-monto y pide el valor limpio al validar y enviar', 'V.valorCampo(campos[k])' in html and 'data-monto' in html)
    html = c.get('/referenciales/item/item-index').get_data(as_text=True)
    chk('Ítems: precio y cantidad con formato, valor limpio al validar y enviar, stock con miles', html.count('V.valorCampo(campos[k])') == 2 and 'V.formatearMonto(f.cantidad_disponible)' in html
        and 'V.formatearMonto(registro.precio_unitario)' in html)
    html = c.get('/referenciales/apertura-caja/apertura-caja-index').get_data(as_text=True)
    chk('Apertura: valida y envía el monto sin puntos', 'V.monto(V.limpiarMonto(txtMonto.value))' in html and 'monto_apertura: V.limpiarMonto(txtMonto.value)' in html)
    html = c.get('/referenciales/arqueo/arqueo-index').get_data(as_text=True)
    chk('Arqueo: guarda lo contado limpio y lo muestra formateado', 'contados[el.dataset.forma] = V.limpiarMonto(el.value)' in html and 'data-monto value="${V.formatearMonto(' in html)
    html = c.get('/referenciales/cobro/cobro-index').get_data(as_text=True)
    chk('Cobro: el monto de cada forma se guarda limpio y se muestra formateado', "l[c] = (c === 'monto') ? V.limpiarMonto(v) : v" in html and 'V.formatearMonto(l.monto)' in html)
    html = c.get('/referenciales/nota-credito-debito/nota-index').get_data(as_text=True)
    chk('Notas: monto y cantidad a devolver limpios al guardar, formateados al mostrar', "V.limpiarMonto(el.value)" in html and 'V.formatearMonto(devuelto[i.id_item]' in html and 'V.formatearMonto(r.monto)' in html)
    html = c.get('/referenciales/pedido-cliente/pedido-cliente-index').get_data(as_text=True)
    chk('Pedido: cantidad limpia al guardar, formateada al mostrar (también el stock)', 'l.cantidad = V.limpiarMonto(inp.value)' in html and 'V.formatearMonto(l.stock)' in html)
    html = c.get('/referenciales/remision/remision-index').get_data(as_text=True)
    chk('Remisión: cantidad limpia al guardar, formateada al mostrar', 'cantidades[el.dataset.item] = V.limpiarMonto(el.value)' in html and 'V.formatearMonto(cantidades[i.id_item]' in html)

    print('\n== PDFs')
    chk('PDF de factura: _miles es el helper único', factura_pdf_service._miles(1500000) == '1.500.000' and factura_pdf_service._miles(0) == '0')
    from app.Services import recibo_pdf_service, remision_pdf_service, arqueo_pdf_service, nota_pdf_service
    chk('recibo, remisión, arqueo y nota importan el mismo helper', all(m._miles is factura_pdf_service._miles for m in (recibo_pdf_service, remision_pdf_service, arqueo_pdf_service))
        and nota_pdf_service._miles is factura_pdf_service._miles)

    print('\n== Informes (pantalla, PDF y Excel)')
    html = c.get('/informes').get_data(as_text=True)
    chk('Informes carga facturacion-validadores.js antes de informes.js (usa formatearMonto)', html.index('facturacion-validadores.js') < html.index('js/informes.js'))
    chk('informes.js formatea con formatearMonto', 'V.formatearMonto' in c.get('/static/js/informes.js').get_data(as_text=True))
    datos = {'kpis': [{'clave': 'total', 'titulo': 'Total', 'valor': 1500000, 'formato': 'gs'}, {'clave': 'n', 'titulo': 'Unidades', 'valor': 12500, 'formato': 'num'},
                      {'clave': 'p', 'titulo': 'Tasa', 'valor': 33.3, 'formato': 'pct'}],
             'tablas': [{'id': 't', 'titulo': 'Detalle', 'columnas': [{'campo': 'a', 'titulo': 'Concepto', 'tipo': 'texto'}, {'campo': 'm', 'titulo': 'Monto', 'tipo': 'gs'},
                                                                      {'campo': 'q', 'titulo': 'Cant.', 'tipo': 'num'}],
                         'filas': [{'a': 'A', 'm': 2500000, 'q': 1500}, {'a': 'B', 'm': 100000, 'q': 3}]}]}
    armado = ex.calcular_datos_exportacion('Prueba', {'razon_social': 'EMPRESA ZZ', 'ruc': '1-1', 'direccion': 'X', 'telefono': '0213398423'}, [('Período', 'p')], datos, 'ZZ')
    txt = texto_pdf(ex.generarInformePdf(armado))
    chk('PDF de informes: "Gs. 1.500.000", "Gs. 2.500.000", "12.500" y "33,3 %"', all(x in txt for x in ('Gs. 1.500.000', 'Gs. 2.500.000', '12.500', '33,3 %')))
    from openpyxl import load_workbook
    libro = load_workbook(io.BytesIO(ex.generarInformeExcel(armado)))
    hoja = libro['Detalle']
    celda_m, celda_q = hoja['B4'], hoja['C4']
    chk('Excel: el monto es un NÚMERO (no texto) con formato de miles #,##0', isinstance(celda_m.value, (int, float)) and celda_m.value == 2500000 and celda_m.number_format == '#,##0')
    chk('Excel: la cantidad también es número con formato de miles', isinstance(celda_q.value, (int, float)) and celda_q.value == 1500 and celda_q.number_format in ('#,##0.##', '#,##0'))
    ind = libro['Indicadores']
    kp = {f[0].value: f[1] for f in ind.iter_rows() if f[0].value}
    chk('Excel: el indicador en guaraníes es número con formato de miles', isinstance(kp['Total'].value, (int, float)) and kp['Total'].value == 1500000 and kp['Total'].number_format == '#,##0')

finally:
    if creados['items']:
        cu.execute("DELETE FROM items WHERE id_item = ANY(%s)", (creados['items'],))
    if creados['tipos_tratamiento']:
        cu.execute("DELETE FROM tipos_tratamiento WHERE id_tipo_tratamiento = ANY(%s)", (creados['tipos_tratamiento'],))
    cu.execute("SELECT id_auditoria_cambio FROM auditoria_cambios WHERE id_auditoria_cambio > %s AND id_usuario = 1 AND ip = '127.0.0.1'", (BASE,))
    ids = [x[0] for x in cu.fetchall()]
    cu.execute("SET session_replication_role = replica")
    cu.execute("DELETE FROM auditoria_cambios WHERE id_auditoria_cambio = ANY(%s)", (ids,))
    cu.execute("RESET session_replication_role")
    print(f'(limpieza: {sum(len(v) for v in creados.values())} registros y {len(ids)} filas de auditoría de prueba borrados por ID)')
    cu.close(); cx.close()

print('FALLOS:', 'ninguno' if not fallos else fallos)
