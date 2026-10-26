"""
=====================================================
Servicio: exportación de los Informes a PDF y Excel
Descripción: arma al vuelo, sin guardar archivos, el PDF (A4 apaisado, mismo estilo que la factura y la auditoría) y el libro de Excel de una
pestaña de Informes a partir del resultado del DAO (`kpis` y `tablas`; los gráficos no se exportan, sus datos están en las tablas y en los
indicadores). Mismo patrón que auditoria_pdf_service.py:
  1. calcular_datos_exportacion(...) -> CÁLCULO: dict puro con todo lo que se imprime (cabecera, filtros, secciones ya formateadas). Sin ReportLab.
  2. dibujar_informe(c, datos)       -> DIBUJO: layout; una sección por tabla, repite su título y encabezado al cambiar de página.
  3. generarInformePdf(datos) / generarInformeExcel(datos) -> bytes del archivo.
=====================================================
"""
import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas

from app.Services.auditoria_pdf_service import (ALTO, DER, INTERLINEADO, MARGEN, TAM, Y_MIN, _cabecera_primera, _encabezado_tabla, _lineas_fila,
                                                calcular_datos_informe)

TOPE_FILAS = 5000     # el DAO corta cada tabla en este tope


# ======================================================================================================================
# 1. CÁLCULO (datos)
# ======================================================================================================================
from app.Services.formato import formatear_gs, formatear_monto as _miles   # montos y cantidades: helper único (app/Services/formato.py)


def formatear(valor, tipo):
    """Texto de un valor para el PDF: guaraníes y enteros con punto de miles, decimales con coma, porcentaje con ' %'."""
    if valor is None or valor == '':
        return ''
    if tipo == 'gs':
        return formatear_gs(valor)
    if tipo == 'pct':
        return f'{valor:.1f}'.replace('.', ',') + ' %'
    if tipo == 'num':
        return _miles(valor) if float(valor).is_integer() else f'{valor:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return str(valor)


def _secciones(kpis, tablas):
    """Secciones del PDF como (título, columnas [(encabezado, ancho relativo, alineación)], filas de texto): indicadores y luego cada tabla."""
    secciones = []
    if kpis:
        secciones.append(('Indicadores', [('Indicador', 3.0, 'L'), ('Valor', 1.5, 'R')], [[k['titulo'], formatear(k['valor'], k['formato'])] for k in kpis]))
    for t in tablas:
        columnas = []
        for c in t['columnas']:
            largo = max([len(str(c['titulo']))] + [len(str(f.get(c['campo'], ''))) for f in t['filas'][:200]])
            columnas.append((c['titulo'], max(1.0, min(largo, 40) * 0.12), 'L' if c['tipo'] == 'texto' else 'R'))
        secciones.append((t['titulo'], columnas, [[formatear(f.get(c['campo']), c['tipo']) for c in t['columnas']] for f in t['filas']]))
    return secciones


def aviso_de_recorte(tablas):
    recortadas = [t['titulo'] for t in tablas if len(t['filas']) >= TOPE_FILAS]
    if not recortadas:
        return ''
    return f'Se muestran las primeras {_miles(TOPE_FILAS)} filas de: ' + ', '.join(recortadas) + '. Acote las fechas para ver el resto.'


def calcular_datos_exportacion(titulo, empresa, filtros, datos, emitido_por, emitido=None):
    """titulo: nombre del informe. filtros: lista de (etiqueta, valor) ya legibles. datos: el dict del DAO. Devuelve lo que imprimen el PDF y el Excel."""
    base = calcular_datos_informe(f'INFORME DE {titulo.upper()}', empresa, filtros, [('', 1, 'L')], [], emitido_por, emitido=emitido,
                                  aviso=aviso_de_recorte(datos['tablas']))
    base['secciones'] = []
    for subtitulo, columnas, filas in _secciones(datos['kpis'], datos['tablas']):
        d = calcular_datos_informe(subtitulo, empresa, filtros, columnas, filas, emitido_por, emitido=emitido)
        base['secciones'].append({'titulo': subtitulo, 'columnas': d['columnas'], 'filas': d['filas']})
    base['informe'] = titulo
    base['crudo'] = datos            # valores sin formatear, para el Excel
    return base


def total_filas(datos):
    return sum(len(t['filas']) for t in datos['tablas'])


# ======================================================================================================================
# 2. DIBUJO (layout)
# ======================================================================================================================
def _pie(c, pagina):
    c.setFont('Helvetica', 7.5)
    c.drawString(MARGEN, MARGEN - 0.2 * cm, 'Informe de solo consulta. Los gráficos de la pantalla se exportan como tablas.')
    c.drawRightString(DER, MARGEN - 0.2 * cm, f'Página {pagina}')


def _titulo_seccion(c, texto, y):
    c.setFillColorRGB(0.92, 0.95, 0.96)
    c.rect(MARGEN, y - 0.2 * cm, DER - MARGEN, 0.6 * cm, stroke=0, fill=1)
    c.setFillColorRGB(0, 0, 0)
    c.setFont('Helvetica-Bold', 9.5)
    c.drawString(MARGEN + 0.15 * cm, y, texto)
    return y - 0.85 * cm


def dibujar_informe(c, d):
    """Dibuja todas las páginas: cabecera en la primera, luego cada sección (indicadores y tablas) con su título y encabezado repetidos."""
    estado = {'pagina': 1}
    y = _cabecera_primera(c, d, ALTO - MARGEN)

    def nueva_pagina(titulo_seccion):
        _pie(c, estado['pagina'])
        c.showPage()
        estado['pagina'] += 1
        c.setFont('Helvetica', 8)
        c.drawString(MARGEN, ALTO - MARGEN, f"{d['titulo']} (continuación)")
        return _titulo_seccion(c, titulo_seccion + ' (continuación)', ALTO - MARGEN - 0.8 * cm)

    for sec in d['secciones']:
        if y - 3.0 * cm < Y_MIN:       # no dejar un título solo al pie de la página
            y = nueva_pagina(sec['titulo'])
        else:
            y = _titulo_seccion(c, sec['titulo'], y)
        tabla = {'columnas': sec['columnas']}
        y = _encabezado_tabla(c, tabla, y)
        if not sec['filas']:
            c.setFont('Helvetica-Oblique', 9)
            c.drawString(MARGEN, y, 'Sin datos para los filtros indicados.')
            y -= 0.7 * cm
        for fila in sec['filas']:
            cols = _lineas_fila(tabla, fila)
            alto = max(len(l) for l in cols) * INTERLINEADO + 0.15 * cm
            if y - alto < Y_MIN:
                y = _encabezado_tabla(c, tabla, nueva_pagina(sec['titulo']))
            c.setFont('Helvetica', TAM)
            x = MARGEN
            for col, lineas in zip(sec['columnas'], cols):
                for i, linea in enumerate(lineas):
                    if col['alinea'] == 'R':
                        c.drawRightString(x + col['ancho'] - 0.1 * cm, y - i * INTERLINEADO, linea)
                    else:
                        c.drawString(x + 0.1 * cm, y - i * INTERLINEADO, linea)
                x += col['ancho']
            y -= alto
            c.setLineWidth(0.25)
            c.line(MARGEN, y + 0.1 * cm, DER, y + 0.1 * cm)
        y -= 0.6 * cm
    _pie(c, estado['pagina'])
    c.showPage()


# ======================================================================================================================
# 3. ARMADO
# ======================================================================================================================
def generarInformePdf(datos):
    """Bytes del PDF (datos = calcular_datos_exportacion)."""
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=landscape(A4))
    c.setTitle(datos['titulo'])
    c.setAuthor(datos['empresa']['razon_social'])
    dibujar_informe(c, datos)
    c.save()
    pdf = buffer.getvalue()
    buffer.close()
    return pdf


def _seguro(valor):
    """Un texto que empieza con = + - @ no se interpreta como fórmula al abrirlo en Excel."""
    v = str(valor)
    return "'" + v if v[:1] in ('=', '+', '-', '@') else v


def _nombre_hoja(titulo, usados):
    base = re.sub(r'[\\/*?:\[\]]', '', titulo)[:28].strip() or 'Hoja'
    nombre, n = base, 2
    while nombre.lower() in usados:
        nombre = f'{base[:25]} {n}'
        n += 1
    usados.add(nombre.lower())
    return nombre


FORMATO_EXCEL = {'gs': '#,##0', 'num': '#,##0.##', 'pct': '0.0"%"'}


def generarInformeExcel(datos):
    """Bytes del .xlsx: una hoja "Indicadores" (con la cabecera, los filtros y los indicadores) y una hoja por tabla, con números reales
    (no texto) para poder sumar y filtrar. Las tablas están cortadas en el mismo tope que la pantalla."""
    crudo = datos['crudo']
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Indicadores'
    negrita = Font(bold=True)
    relleno = PatternFill('solid', fgColor='E6F0F2')
    hoja.append([datos['titulo']]); hoja['A1'].font = Font(bold=True, size=14)
    e = datos['empresa']
    hoja.append([f"{e['razon_social']}  RUC: {e['ruc']}"])
    hoja.append(['Emitido', datos['emitido'], 'Por', datos['emitido_por']])
    for k, v in datos['filtros']:
        hoja.append(['Filtro', k, v])
    if datos['aviso']:
        hoja.append(['Aviso', datos['aviso']])
    hoja.append([])
    if crudo['kpis']:
        hoja.append(['Indicador', 'Valor'])
        for celda in hoja[hoja.max_row]:
            celda.font = negrita; celda.fill = relleno
        for k in crudo['kpis']:
            hoja.append([k['titulo'], k['valor']])
            celda = hoja.cell(row=hoja.max_row, column=2)
            celda.number_format = FORMATO_EXCEL.get(k['formato'], 'General')
            celda.alignment = Alignment(horizontal='right')
    hoja.column_dimensions['A'].width = 38
    hoja.column_dimensions['B'].width = 24
    hoja.column_dimensions['C'].width = 24
    usados = {'indicadores'}
    for t in crudo['tablas']:
        h = libro.create_sheet(_nombre_hoja(t['titulo'], usados))
        h.append([t['titulo']]); h['A1'].font = Font(bold=True, size=12)
        h.append([])
        h.append([c['titulo'] for c in t['columnas']])
        for celda in h[3]:
            celda.font = negrita; celda.fill = relleno
        for f in t['filas']:
            h.append([_seguro(f.get(c['campo'])) if c['tipo'] == 'texto' and f.get(c['campo']) is not None else f.get(c['campo']) for c in t['columnas']])
        for j, c in enumerate(t['columnas'], 1):
            ancho = max([len(str(c['titulo']))] + [len(str(f.get(c['campo'], ''))) for f in t['filas'][:300]])
            h.column_dimensions[get_column_letter(j)].width = min(max(ancho + 2, 10), 50)
            if c['tipo'] != 'texto':
                for fila in range(4, h.max_row + 1):
                    celda = h.cell(row=fila, column=j)
                    celda.number_format = FORMATO_EXCEL.get(c['tipo'], 'General')
        h.freeze_panes = 'A4'
        if t['filas']:
            h.auto_filter.ref = f'A3:{get_column_letter(len(t["columnas"]))}{h.max_row}'
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
